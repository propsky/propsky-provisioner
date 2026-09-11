from __future__ import annotations

import csv
import shutil
import threading
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from .models import LEDGER_HEADERS, LedgerRecord


class Ledger:
    """Write the CSV first, then best-effort synchronize the operator workbook."""

    def __init__(self, root: Path) -> None:
        self.root = root
        self.csv_path = root / "provision_log.csv"
        self.xlsx_path = root / "出廠帳本.xlsx"
        self._lock = threading.RLock()

    def backup(self) -> Path | None:
        if not self.xlsx_path.exists():
            return None
        backup = self.xlsx_path.with_name(
            f"{self.xlsx_path.stem}.backup-{datetime.now():%Y%m%d-%H%M%S}{self.xlsx_path.suffix}"
        )
        shutil.copy2(self.xlsx_path, backup)
        return backup

    def existing_records(self) -> list[dict[str, str]]:
        with self._lock:
            rows: list[dict[str, str]] = []
            if self.csv_path.exists():
                with self.csv_path.open("r", encoding="utf-8-sig", newline="") as stream:
                    rows.extend(dict(row) for row in csv.DictReader(stream))
            return rows

    def duplicate_fields(self) -> list[str]:
        """Return identity fields whose mappings conflict, not repeated history rows."""
        with self._lock:
            rows = self.existing_records()
            if self.xlsx_path.exists():
                workbook = load_workbook(self.xlsx_path, read_only=True, data_only=True)
                try:
                    sheet = workbook.active
                    headers = [str(cell.value or "") for cell in sheet[1]]
                    rows.extend(
                        {
                            header: str(values[index] or "")
                            for index, header in enumerate(headers)
                            if header and index < len(values)
                        }
                        for values in sheet.iter_rows(min_row=2, values_only=True)
                    )
                finally:
                    workbook.close()

            conflicts: list[str] = []
            mac_mappings: dict[str, set[tuple[str, str]]] = {}
            card_mappings: dict[str, set[str]] = {}
            token_mappings: dict[str, set[str]] = {}
            for row in rows:
                mac = row.get("MAC Address (CPUID)", "").strip().upper()
                card = row.get("小卡編號", "").strip().upper()
                token = row.get("UUID (token)", "").strip().upper()
                if mac:
                    mac_mappings.setdefault(mac, set()).add((card, token))
                if card and mac:
                    card_mappings.setdefault(card, set()).add(mac)
                if token and mac:
                    token_mappings.setdefault(token, set()).add(mac)
            if any(len(values) > 1 for values in mac_mappings.values()):
                conflicts.append("MAC Address (CPUID)")
            if any(len(values) > 1 for values in card_mappings.values()):
                conflicts.append("小卡編號")
            if any(len(values) > 1 for values in token_mappings.values()):
                conflicts.append("UUID (token)")
            return conflicts

    def find_by_mac(self, mac: str) -> dict[str, str] | None:
        with self._lock:
            normalized = mac.upper()
            rows = self.existing_records()
            for row in reversed(rows):
                if row.get("MAC Address (CPUID)", "").upper() == normalized:
                    return row
            if not self.xlsx_path.exists():
                return None
            workbook = load_workbook(self.xlsx_path, read_only=True, data_only=True)
            try:
                sheet = workbook.active
                headers = [str(cell.value or "") for cell in sheet[1]]
                for values in reversed(list(sheet.iter_rows(min_row=2, values_only=True))):
                    row = {header: str(values[index] or "") for index, header in enumerate(headers) if header}
                    if row.get("MAC Address (CPUID)", "").upper() == normalized:
                        return row
                return None
            finally:
                workbook.close()

    def append_csv(self, record: LedgerRecord) -> None:
        with self._lock:
            self.root.mkdir(parents=True, exist_ok=True)
            new_file = not self.csv_path.exists() or self.csv_path.stat().st_size == 0
            with self.csv_path.open("a", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=LEDGER_HEADERS)
                if new_file:
                    writer.writeheader()
                writer.writerow(self._row(record))

    def sync_xlsx(self, record: LedgerRecord) -> None:
        with self._lock:
            self.root.mkdir(parents=True, exist_ok=True)
            if self.xlsx_path.exists():
                workbook = load_workbook(self.xlsx_path)
                sheet = workbook.active
                current_headers = [str(cell.value or "") for cell in sheet[1]]
                if current_headers != list(LEDGER_HEADERS):
                    old_rows = list(sheet.iter_rows(min_row=2, values_only=True))
                    workbook.remove(sheet)
                    sheet = workbook.create_sheet("出廠紀錄", 0)
                    sheet.append(LEDGER_HEADERS)
                    for values in old_rows:
                        old = {header: values[index] for index, header in enumerate(current_headers) if index < len(values)}
                        sheet.append([old.get(header, "") for header in LEDGER_HEADERS])
            else:
                workbook = Workbook()
                sheet = workbook.active
                sheet.title = "出廠紀錄"
                sheet.append(LEDGER_HEADERS)

            row = self._row(record)
            mac_column = LEDGER_HEADERS.index("MAC Address (CPUID)") + 1
            target_row = None
            for row_number in range(sheet.max_row, 1, -1):
                if str(sheet.cell(row_number, mac_column).value or "").upper() == (record.mac or "").upper():
                    target_row = row_number
                    break
            if target_row is None:
                target_row = sheet.max_row + 1
            for column, header in enumerate(LEDGER_HEADERS, 1):
                sheet.cell(target_row, column).value = row[header]
            workbook.save(self.xlsx_path)

    def update_result(self, mac: str, result: str, note: str = "") -> None:
        with self._lock:
            rows = self.existing_records()
            target = next((row for row in reversed(rows) if row.get("MAC Address (CPUID)", "").upper() == mac.upper()), None)
            if target is None:
                return
            target["結果"] = result
            if note:
                target["備註"] = note
            with self.csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=LEDGER_HEADERS)
                writer.writeheader()
                writer.writerows(rows)
            if not self.xlsx_path.exists():
                return
            workbook = load_workbook(self.xlsx_path)
            sheet = workbook.active
            mac_column = LEDGER_HEADERS.index("MAC Address (CPUID)") + 1
            result_column = LEDGER_HEADERS.index("結果") + 1
            note_column = LEDGER_HEADERS.index("備註") + 1
            for row_number in range(sheet.max_row, 1, -1):
                if str(sheet.cell(row_number, mac_column).value or "").upper() == mac.upper():
                    sheet.cell(row_number, result_column).value = result
                    if note:
                        sheet.cell(row_number, note_column).value = note
                    break
            workbook.save(self.xlsx_path)

    def update_details(self, mac: str, **details: str) -> None:
        with self._lock:
            rows = self.existing_records()
            target = next((row for row in reversed(rows) if row.get("MAC Address (CPUID)", "").upper() == mac.upper()), None)
            if target is None:
                return
            columns = {"app_version": "程式版本", "rssi": "RSSI", "mqtt": "MQTT"}
            target.update({columns[key]: value for key, value in details.items() if key in columns})
            with self.csv_path.open("w", encoding="utf-8-sig", newline="") as stream:
                writer = csv.DictWriter(stream, fieldnames=LEDGER_HEADERS)
                writer.writeheader()
                writer.writerows(rows)
            if self.xlsx_path.exists():
                workbook = load_workbook(self.xlsx_path)
                sheet = workbook.active
                mac_column = LEDGER_HEADERS.index("MAC Address (CPUID)") + 1
                for row_number in range(sheet.max_row, 1, -1):
                    if str(sheet.cell(row_number, mac_column).value or "").upper() == mac.upper():
                        for key, value in details.items():
                            column_name = columns.get(key)
                            if column_name:
                                sheet.cell(row_number, LEDGER_HEADERS.index(column_name) + 1).value = value
                        break
                workbook.save(self.xlsx_path)

    @staticmethod
    def _row(record: LedgerRecord) -> dict[str, str]:
        # Password is intentionally not part of LedgerRecord and can never reach disk.
        return {
            "小卡編號": record.card_number or "",
            "MAC Address (CPUID)": record.mac or "",
            "UUID (token)": record.token or "",
            "燒錄時間": record.burned_at.isoformat(timespec="seconds"),
            "Port": record.port,
            "MicroPython": record.micropython or "",
            "程式版本": record.app_version or "",
            "SSID": record.ssid or "",
            "RSSI": record.rssi or "",
            "MQTT": record.mqtt or "",
            "結果": record.result,
            "備註": record.note,
        }
